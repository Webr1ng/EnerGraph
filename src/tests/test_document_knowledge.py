"""test_document_knowledge — 上传文件 RAG 的生命周期、检索和路由回归

所属层：tests
依赖：pytest, langchain_core, src.services.document_knowledge
对接算法层：N/A（Mock ChromaDB）
"""
from pathlib import Path

import pytest
from langchain_core.messages import AIMessage

from src.graph.nodes import _enforce_hvac_tool_route
from src.services.document_knowledge import DocumentKnowledgeService, DocumentProcessingError
from src.skills.document_knowledge_skill import DocumentKnowledgeSkill
from src.tools import TOOL_REGISTRY


class FakeCollection:
    """仅实现文档服务所需 Chroma collection 契约的内存替身。"""

    def __init__(self) -> None:
        self.items: dict[str, tuple[str, dict]] = {}

    def upsert(self, ids, documents, metadatas) -> None:
        for identifier, document, metadata in zip(ids, documents, metadatas):
            self.items[identifier] = (document, metadata)

    def delete(self, where) -> None:
        self.items = {
            identifier: item
            for identifier, item in self.items.items()
            if item[1].get("document_id") != where.get("document_id")
        }

    def count(self) -> int:
        return len(self.items)

    def query(self, query_texts, n_results, include) -> dict:
        selected = list(self.items.values())[:n_results]
        return {
            "documents": [[item[0] for item in selected]],
            "distances": [[0.2 for _ in selected]],
            "metadatas": [[item[1] for item in selected]],
        }


def _write_text_pdf(path: Path) -> None:
    """生成最小文字型 PDF，用于验证 pypdf 的真实提取路径。"""
    content = b"BT\n/F1 12 Tf\n72 720 Td\n(PDF maintenance procedure) Tj\nET\n"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 5 0 R >> >> /MediaBox [0 0 612 792] /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode())
        data.extend(body)
        data.extend(b"\nendobj\n")
    xref_offset = len(data)
    data.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    data.extend(b"0000000000 65535 f \n")
    data.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    data.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    )
    path.write_bytes(bytes(data))

@pytest.fixture
def service(tmp_path: Path, monkeypatch) -> DocumentKnowledgeService:
    """创建不依赖本地 embedding 或真实 Chroma 的文档服务。"""
    instance = DocumentKnowledgeService(storage_dir=tmp_path / "uploads")
    collection = FakeCollection()
    monkeypatch.setattr(instance, "_get_collection", lambda create: collection)
    return instance


def test_txt_upload_recall_duplicate_and_delete(service: DocumentKnowledgeService) -> None:
    """TXT 入库后应可引用召回；重复上传复用；删除后不能再检索。"""
    result = service.upload_bytes("运行手册.txt", "冷水机组应每月检查冷却水温度。".encode())

    assert result.document.status.value == "ready"
    assert result.document.chunk_count == 1
    recalled = service.search("冷却水温度怎么检查")
    assert recalled.low_confidence is False
    assert recalled.sources[0].file_name == "运行手册.txt"

    duplicate = service.upload_bytes("副本.txt", "冷水机组应每月检查冷却水温度。".encode())
    assert duplicate.duplicate is True
    assert duplicate.document.document_id == result.document.document_id

    assert service.delete_document(result.document.document_id) is True
    assert service.search("冷却水温度怎么检查").low_confidence is True


def test_json_upload_keeps_field_content(service: DocumentKnowledgeService) -> None:
    """JSON 应被展开成可检索字段路径和值。"""
    result = service.upload_bytes(
        "设备.json", '{"device":{"name":"冷水机","interval":"每月"}}'.encode()
    )

    assert result.document.status.value == "ready"
    recalled = service.search("设备检查周期")
    assert "device.name: 冷水机" in recalled.results[0]


def test_docx_and_text_pdf_extract_real_content(service: DocumentKnowledgeService, tmp_path: Path) -> None:
    """DOCX 和文字型 PDF 解析必须保留章节或页码及正文内容。"""
    from docx import Document

    docx_path = tmp_path / "manual.docx"
    docx = Document()
    docx.add_heading("维护要求", level=1)
    docx.add_paragraph("每月检查冷却水温度。")
    docx.save(docx_path)
    docx_sections = service._parse_docx(docx_path)
    assert docx_sections[0][0] == "维护要求"
    assert "冷却水温度" in docx_sections[0][2]

    pdf_path = tmp_path / "manual.pdf"
    _write_text_pdf(pdf_path)
    pdf_sections = service._parse_pdf(pdf_path)
    assert pdf_sections[0][1] == 1
    assert "PDF maintenance procedure" in pdf_sections[0][2]


def test_scanned_pdf_is_deferred_for_ocr(service: DocumentKnowledgeService, tmp_path: Path) -> None:
    """没有文字层的 PDF 必须失败并提示 OCR 后续支持。"""
    from pypdf import PdfWriter

    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as stream:
        writer.write(stream)
    with pytest.raises(DocumentProcessingError, match="OCR"):
        service._parse_pdf(path)


@pytest.mark.parametrize("suffix", [".doc", ".docx", ".pdf"])
def test_supported_binary_formats_enter_parse_pipeline(
    service: DocumentKnowledgeService, monkeypatch, suffix: str
) -> None:
    """DOC/DOCX/PDF 均应进入统一解析、切块、入库主流程。"""
    captured: list[str] = []

    def parse_file(_path: Path, file_type: str):
        captured.append(file_type)
        return [("测试章节", 2 if suffix == ".pdf" else None, "测试文档正文")]

    monkeypatch.setattr(service, "_parse_file", parse_file)
    result = service.upload_bytes(f"资料{suffix}", b"not-a-real-document")

    assert result.document.status.value == "ready"
    assert captured == [suffix.lstrip(".")]


def test_empty_unsupported_and_scanned_pdf_failure(service: DocumentKnowledgeService, monkeypatch) -> None:
    """空文件、不支持格式和无法抽取文字的 PDF 必须有确定结果。"""
    with pytest.raises(DocumentProcessingError, match="文件为空"):
        service.upload_bytes("空.txt", b"")
    with pytest.raises(DocumentProcessingError, match="不支持"):
        service.upload_bytes("表格.xlsx", b"data")

    monkeypatch.setattr(
        service,
        "_parse_file",
        lambda _path, _kind: (_ for _ in ()).throw(
            DocumentProcessingError("PDF 未提取到文字，扫描件 OCR 将在后续版本支持")
        ),
    )
    result = service.upload_bytes("扫描件.pdf", b"pdf")
    assert result.document.status.value == "failed"
    assert "OCR" in (result.document.failure_reason or "")


def test_reprocess_updates_retry_count(service: DocumentKnowledgeService, monkeypatch) -> None:
    """重新解析必须从原文件执行并累加重试次数。"""
    monkeypatch.setattr(service, "_parse_file", lambda *_args: [("正文", None, "初始内容")])
    uploaded = service.upload_bytes("重试.txt", b"content")
    monkeypatch.setattr(service, "_parse_file", lambda *_args: [("正文", None, "重新解析内容")])

    result = service.reprocess(uploaded.document.document_id)
    assert result.document.status.value == "ready"
    assert result.document.retry_count == 1
    assert "重新解析内容" in service.search("测试").results[0]


def test_delete_keeps_document_when_vector_cleanup_fails(
    service: DocumentKnowledgeService, monkeypatch
) -> None:
    """向量删除失败时不能孤立删除登记和原文件。"""
    uploaded = service.upload_bytes("一致性.txt", "需要同步删除".encode())
    collection = service._get_collection(create=False)
    monkeypatch.setattr(
        collection,
        "delete",
        lambda **_kwargs: (_ for _ in ()).throw(RuntimeError("Chroma unavailable")),
    )

    with pytest.raises(DocumentProcessingError, match="无法删除文档向量"):
        service.delete_document(uploaded.document.document_id)
    assert service.get_document(uploaded.document.document_id) is not None


def test_document_route_is_forced_without_llm_tool_call() -> None:
    """明确依据上传资料提问时，本地模型漏调 Tool 也必须补上文档 RAG。"""
    response = AIMessage(content="我来回答")
    routed = _enforce_hvac_tool_route(
        response,
        "根据我上传的运维手册，冷水机怎么维护？",
        {"user_input": "根据我上传的运维手册，冷水机怎么维护？"},
    )

    assert "query_uploaded_documents" in [call["name"] for call in routed.tool_calls]
    assert "query_uploaded_documents" in TOOL_REGISTRY


def test_document_skill_refuses_low_confidence_and_formats_source(monkeypatch) -> None:
    """文档 Skill 必须在低置信度拒答，并为命中结果注入文件级引用。"""
    monkeypatch.setattr(
        "src.skills.document_knowledge_skill.load_prompts",
        lambda: {
            "document_knowledge_refusal": {"system": "未在知识库找到可靠内容"},
            "document_knowledge_citation_format": {"system": "必须列出文件页码"},
        },
    )
    skill = DocumentKnowledgeSkill()
    refused = skill.execute(
        [("query_uploaded_documents", {"query": "问题", "low_confidence": True}, {})], {}
    )
    assert refused["document_context_hint"]["low_confidence"] is True
    assert "未在知识库" in refused["document_context_hint"]["system_suffix"]

    cited = skill.execute(
        [("query_uploaded_documents", {
            "query": "问题", "low_confidence": False, "results": ["片段"],
            "sources": [{"file_name": "手册.pdf", "page": 3, "section": "维护"}],
        }, {})], {}
    )
    assert "手册.pdf，第 3 页，维护" in cited["document_context_hint"]["system_suffix"]
